from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io
import re

def markdown_to_docx(markdown_text):
    """
    Converts a Markdown string to a professional DOCX byte stream.
    Supports basic headings, bold text, and lists.
    """
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    lines = markdown_text.split('\n')
    for line in lines:
        stripped_line = line.strip()
        if not stripped_line:
            doc.add_paragraph()
            continue
            
        # 1. Headings
        header_match = re.match(r'^(#{1,6})\s*(.*)', stripped_line)
        if header_match:
            level = len(header_match.group(1))
            text = header_match.group(2)
            # Remove any internal markdown from header text
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            text = re.sub(r'\*(.*?)\*', r'\1', text)
            doc.add_heading(text, level=level)
            continue

        # 2. Lists
        list_match = re.match(r'^[\*\-\+]\s+(.*)', stripped_line)
        if list_match:
            text = list_match.group(1)
            p = doc.add_paragraph(style='List Bullet')
            apply_formatting(p, text)
            continue
            
        number_list_match = re.match(r'^\d+\.\s+(.*)', stripped_line)
        if number_list_match:
            text = number_list_match.group(1)
            p = doc.add_paragraph(style='List Number')
            apply_formatting(p, text)
            continue

        # 3. Normal paragraph
        p = doc.add_paragraph()
        apply_formatting(p, stripped_line)
        
    # Save to bytes
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def apply_formatting(paragraph, text):
    """Parses bold and italic markdown in a line of text."""
    # Split by bold markers first
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            inner_text = part[2:-2]
            # Handle italics within bold
            sub_parts = re.split(r'(\*.*?\*)', inner_text)
            for sub in sub_parts:
                if sub.startswith('*') and sub.endswith('*'):
                    run = paragraph.add_run(sub[1:-1])
                    run.bold = True
                    run.italic = True
                else:
                    run = paragraph.add_run(sub)
                    run.bold = True
        else:
            # Handle italics
            sub_parts = re.split(r'(\*.*?\*)', part)
            for sub in sub_parts:
                if sub.startswith('*') and sub.endswith('*'):
                    run = paragraph.add_run(sub[1:-1])
                    run.italic = True
                else:
                    paragraph.add_run(sub)
